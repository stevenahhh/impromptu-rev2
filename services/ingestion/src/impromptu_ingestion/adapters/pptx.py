"""python-pptx structural extraction without rendering or notes export."""

import hashlib
from collections.abc import Iterable, Sequence
from typing import Protocol, cast

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.shapes.base import BaseShape
from pptx.slide import Slide

from impromptu_ingestion.adapters.base import StructuralAdapter, StructuralExtractionError
from impromptu_ingestion.canonical import deck_id, slide_key
from impromptu_ingestion.contracts import (
    ChartElement,
    ChartSeries,
    DeckManifest,
    ExtractionWarning,
    ImageElement,
    InputKind,
    RenderBoundary,
    SlideManifest,
    StructuralElement,
    TableElement,
    TextElement,
    ValidatedInput,
)

_EMU_PER_POINT = 12_700


class _TableCell(Protocol):
    text: str


class _TableRow(Protocol):
    cells: Sequence[_TableCell]


class _Table(Protocol):
    rows: Sequence[_TableRow]


class _TableShape(Protocol):
    table: _Table


class _Series(Protocol):
    name: object | None
    values: Iterable[object | None]


class _Chart(Protocol):
    chart_type: object
    series: Iterable[_Series]


class _ChartShape(Protocol):
    chart: _Chart


class _Image(Protocol):
    blob: bytes
    content_type: str
    size: tuple[int, int]


class _ImageShape(Protocol):
    image: _Image


class _TextFrame(Protocol):
    text: str


class _TextShape(Protocol):
    text_frame: _TextFrame


def _points(value: int) -> float:
    return round(value / _EMU_PER_POINT, 4)


def _position(shape: BaseShape) -> tuple[float, float, float, float]:
    return (
        _points(shape.left),
        _points(shape.top),
        _points(shape.width),
        _points(shape.height),
    )


def _chart_series(shape: _ChartShape) -> tuple[ChartSeries, ...]:
    return tuple(
        ChartSeries(
            name=str(series.name or ""),
            values=tuple("" if value is None else str(value) for value in series.values),
        )
        for series in shape.chart.series
    )


def _extract_shape(shape: BaseShape) -> StructuralElement | None:
    element_id = f"shape:{shape.shape_id}"
    x, y, width, height = _position(shape)

    if shape.has_table:
        table_shape = cast(_TableShape, shape)
        rows = tuple(tuple(cell.text for cell in row.cells) for row in table_shape.table.rows)
        return TableElement(
            element_id=element_id,
            rows=rows,
            x=x,
            y=y,
            width=width,
            height=height,
        )

    if shape.has_chart:
        chart_shape = cast(_ChartShape, shape)
        return ChartElement(
            element_id=element_id,
            chart_type=str(chart_shape.chart.chart_type),
            series=_chart_series(chart_shape),
            x=x,
            y=y,
            width=width,
            height=height,
        )

    if shape.shape_type is MSO_SHAPE_TYPE.PICTURE:
        try:
            image = cast(_ImageShape, shape).image
        except ValueError:
            # A picture linked to an external file carries no embedded bytes, so there is nothing
            # to hash or measure; python-pptx raises rather than returning None. Linked pictures
            # are ordinary in decks authored against a shared drive, and one of them must not cost
            # the whole deck. Fall through so the caller records an unsupported-shape warning and
            # keeps every other element on the slide.
            return None
        pixel_width, pixel_height = image.size
        return ImageElement(
            element_id=element_id,
            content_sha256=hashlib.sha256(image.blob).hexdigest(),
            media_type=image.content_type,
            pixel_width=pixel_width,
            pixel_height=pixel_height,
            x=x,
            y=y,
            width=width,
            height=height,
        )

    if shape.has_text_frame:
        text = cast(_TextShape, shape).text_frame.text.strip()
        if text:
            return TextElement(
                element_id=element_id,
                text=text,
                x=x,
                y=y,
                width=width,
                height=height,
            )
    return None


def _slide_manifest(
    slide: Slide,
    index: int,
    source_sha256: str,
    width_points: float,
    height_points: float,
) -> SlideManifest:
    elements: list[StructuralElement] = []
    warnings: list[ExtractionWarning] = []
    for shape in slide.shapes:
        element = _extract_shape(shape)
        if element is not None:
            elements.append(element)
        elif shape.shape_type not in {MSO_SHAPE_TYPE.PLACEHOLDER, MSO_SHAPE_TYPE.AUTO_SHAPE}:
            warnings.append(
                ExtractionWarning(
                    code="unsupported_shape",
                    message=f"shape {shape.shape_id} ({shape.shape_type}) was not extracted",
                )
            )

    source_id = f"slide:{slide.slide_id}"
    return SlideManifest(
        slide_key=slide_key(source_sha256, source_id),
        source_index=index,
        source_id=source_id,
        width_points=width_points,
        height_points=height_points,
        elements=tuple(elements),
        warnings=tuple(warnings),
    )


class PptxStructuralAdapter(StructuralAdapter):
    """Extract public slide structure; speaker notes are deliberately never read."""

    kind = InputKind.PPTX
    adapter_version = "python-pptx-structural-v1"

    def extract(self, source: ValidatedInput) -> DeckManifest:
        self._require_kind(source)
        try:
            presentation = Presentation(str(source.path))
            slide_width = cast(int, presentation.slide_width)
            slide_height = cast(int, presentation.slide_height)
            slides: Iterable[Slide] = presentation.slides
            manifests = tuple(
                _slide_manifest(
                    slide,
                    index,
                    source.source_sha256,
                    _points(slide_width),
                    _points(slide_height),
                )
                for index, slide in enumerate(slides, start=1)
            )
        except Exception as error:
            # The code stays stable for callers, but the cause has to survive: this blanket catch
            # previously reduced every structural failure to four identical words, which meant a
            # rejected upload could only be diagnosed by re-running the extraction by hand.
            raise StructuralExtractionError(
                "invalid_document",
                f"PPTX structure could not be parsed: {type(error).__name__}: {error}",
            ) from error

        if not manifests:
            raise StructuralExtractionError("empty_document", "PPTX contains no slides")
        return DeckManifest(
            deck_id=deck_id(source.source_sha256),
            source_sha256=source.source_sha256,
            source_kind=self.kind,
            adapter_version=self.adapter_version,
            slides=manifests,
            render_boundary=RenderBoundary.structural_only(),
        )
