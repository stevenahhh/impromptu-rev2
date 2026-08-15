from collections.abc import Mapping

from pptx import Presentation
from pptx.oxml import parse_xml
from pptx.util import Inches

from impromptu_ingestion.render.contracts import (
    ColorBehavior,
    EffectClass,
    EffectTrigger,
    FadeBehavior,
    MotionBehavior,
    ResolvedTarget,
)
from impromptu_ingestion.render.timing import parse_slide_timeline

_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
_SLIDE_KEY = "slide_" + "a" * 64


def _slide_with_shapes(count: int = 2) -> tuple[object, tuple[int, ...]]:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    ids = tuple(
        slide.shapes.add_textbox(Inches(index), Inches(1), Inches(1), Inches(1)).shape_id
        for index in range(1, count + 1)
    )
    return slide, ids


def _targets(ids: tuple[int, ...]) -> Mapping[int, ResolvedTarget]:
    return {
        shape_id: ResolvedTarget(
            shape_id=shape_id,
            shape_name=f"Shape {shape_id}",
            svg_element_id=f"shape-{shape_id}",
        )
        for shape_id in ids
    }


def _effect(
    *,
    preset_id: int,
    preset_class: str,
    node_type: str,
    spid: int,
    behavior: str,
    duration: str = "500",
    text_target: bool = False,
) -> str:
    text_xml = "<p:txtEl><p:charRg st='0' end='2'/></p:txtEl>" if text_target else ""
    return f"""
      <p:par><p:cTn presetID="{preset_id}" presetClass="{preset_class}"
          presetSubtype="0" nodeType="{node_type}">
        <p:stCondLst><p:cond delay="0"/></p:stCondLst>
        <p:childTnLst>{behavior.format(duration=duration, spid=spid, text=text_xml)}</p:childTnLst>
      </p:cTn></p:par>"""


def _timing(*effects: str) -> str:
    return f"""<p:timing xmlns:p="{_P}"><p:tnLst><p:par>
      <p:cTn nodeType="tmRoot"><p:childTnLst><p:seq>
        <p:cTn nodeType="mainSeq"><p:childTnLst><p:par><p:cTn>
          <p:childTnLst>{"".join(effects)}</p:childTnLst>
        </p:cTn></p:par></p:childTnLst></p:cTn>
      </p:seq></p:childTnLst></p:cTn>
    </p:par></p:tnLst></p:timing>"""


def _xml(slide: object, timing: str | None = None, transition: str | None = None) -> bytes:
    element = slide._element  # type: ignore[attr-defined]
    if transition is not None:
        element.append(parse_xml(transition))
    if timing is not None:
        element.append(parse_xml(timing))
    return element.xml.encode()


_FADE = """<p:animEffect transition="in" filter="fade">
  <p:cBhvr><p:cTn dur="{duration}"/>
  <p:tgtEl><p:spTgt spid="{spid}">{text}</p:spTgt></p:tgtEl></p:cBhvr>
</p:animEffect>"""
_MOTION = """<p:animMotion path="M 0 1 L 0 0 L 0 0 E">
  <p:cBhvr><p:cTn dur="{duration}"/>
  <p:tgtEl><p:spTgt spid="{spid}">{text}</p:spTgt></p:tgtEl></p:cBhvr>
</p:animMotion>"""
_COLOR = """<p:anim clrSpc="rgb" from="#0070C0" to="#FF0000">
  <p:cBhvr><p:cTn dur="{duration}"/>
  <p:tgtEl><p:spTgt spid="{spid}">{text}</p:spTgt></p:tgtEl></p:cBhvr>
</p:anim>"""


def test_parses_reference_three_effect_click_group() -> None:
    slide, ids = _slide_with_shapes()
    timing = _timing(
        _effect(
            preset_id=1,
            preset_class="entr",
            node_type="clickEffect",
            spid=ids[0],
            behavior=_FADE,
        ),
        _effect(
            preset_id=2,
            preset_class="path",
            node_type="withEffect",
            spid=ids[1],
            behavior=_MOTION,
            duration="1000",
        ),
        _effect(
            preset_id=11,
            preset_class="emph",
            node_type="afterEffect",
            spid=ids[0],
            behavior=_COLOR,
            duration="300",
        ),
    )

    timeline = parse_slide_timeline(_xml(slide, timing), _SLIDE_KEY, _targets(ids))

    assert len(timeline.click_groups) == 1
    first, second, third = timeline.click_groups[0]
    assert (first.trigger, first.effect_class, first.duration_ms) == (
        EffectTrigger.ON_CLICK,
        EffectClass.ENTRANCE,
        500,
    )
    assert isinstance(first.behavior, FadeBehavior)
    assert first.behavior.direction == "in"
    assert second.trigger is EffectTrigger.WITH_PREVIOUS
    assert isinstance(second.behavior, MotionBehavior)
    assert second.behavior.path == "M 0 1 L 0 0 L 0 0 E"
    assert third.trigger is EffectTrigger.AFTER_PREVIOUS
    assert isinstance(third.behavior, ColorBehavior)
    assert third.behavior.to_color == "#FF0000"
    assert timeline.unsupported == ()


def test_slide_without_timing_has_an_empty_timeline() -> None:
    slide, ids = _slide_with_shapes(1)

    timeline = parse_slide_timeline(_xml(slide), _SLIDE_KEY, _targets(ids))

    assert timeline.click_groups == ()
    assert timeline.unsupported == ()


def test_unknown_preset_class_is_recorded() -> None:
    slide, ids = _slide_with_shapes(1)
    timing = _timing(
        _effect(
            preset_id=99,
            preset_class="mystery",
            node_type="clickEffect",
            spid=ids[0],
            behavior=_FADE,
        )
    )

    timeline = parse_slide_timeline(_xml(slide, timing), _SLIDE_KEY, _targets(ids))

    assert timeline.click_groups == ((),)
    assert timeline.unsupported[0].reason_code == "unknown_preset_class"
    assert timeline.unsupported[0].preset_id == 99


def test_text_range_target_is_recorded() -> None:
    slide, ids = _slide_with_shapes(1)
    timing = _timing(
        _effect(
            preset_id=1,
            preset_class="entr",
            node_type="clickEffect",
            spid=ids[0],
            behavior=_FADE,
            text_target=True,
        )
    )

    timeline = parse_slide_timeline(_xml(slide, timing), _SLIDE_KEY, _targets(ids))

    assert timeline.unsupported[0].reason_code == "text_range_target"


def test_unresolved_shape_id_is_recorded() -> None:
    slide, ids = _slide_with_shapes(1)
    timing = _timing(
        _effect(
            preset_id=1,
            preset_class="entr",
            node_type="clickEffect",
            spid=ids[0] + 100,
            behavior=_FADE,
        )
    )

    timeline = parse_slide_timeline(_xml(slide, timing), _SLIDE_KEY, _targets(ids))

    assert timeline.unsupported[0].reason_code == "unresolved_target"


def test_parses_slide_transition() -> None:
    slide, ids = _slide_with_shapes(1)
    transition = f'<p:transition xmlns:p="{_P}" advClick="0" advTm="2500"><p:fade/></p:transition>'

    timeline = parse_slide_timeline(_xml(slide, transition=transition), _SLIDE_KEY, _targets(ids))

    assert timeline.transition is not None
    assert timeline.transition.kind == "fade"
    assert timeline.transition.advance_on_click is False
