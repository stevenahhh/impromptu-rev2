import json
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from impromptu_ingestion.cli import main


@pytest.fixture
def simple_deck(tmp_path: Path) -> Path:
    path = tmp_path / "발표.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(1), Inches(3), Inches(2)
    )
    shape.text_frame.text = "한국어"
    presentation.save(path)
    return path


def _requires_libreoffice() -> bool:
    from impromptu_ingestion.render.libreoffice import discover_soffice

    return discover_soffice() is None


def test_render_reports_the_missing_renderer_instead_of_pretending(
    simple_deck: Path,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    monkeypatch.setattr("impromptu_ingestion.cli.discover_soffice", lambda: None)

    exit_code = main(["render", str(simple_deck), "--output-dir", str(tmp_path / "out")])

    assert exit_code == 2
    assert "renderer_not_configured" in capsys.readouterr().err


def test_render_publishes_static_png_pages_for_pdf_without_libreoffice(
    sample_pdf: Path,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    output_dir = tmp_path / "pdf-out"
    monkeypatch.setattr("impromptu_ingestion.cli.discover_soffice", lambda: None)

    exit_code = main(["render", str(sample_pdf), "--output-dir", str(output_dir)])

    assert exit_code == 0
    manifest = json.loads((output_dir / "render.json").read_text(encoding="utf-8"))
    assert manifest["animation_eligible"] is False
    assert len(manifest["slides"]) == 2
    assert all(slide["relative_path"].endswith(".png") for slide in manifest["slides"])
    assert all((output_dir / slide["relative_path"]).is_file() for slide in manifest["slides"])
    assert "rendered 2 static PDF pages" in capsys.readouterr().out


@pytest.mark.skipif(_requires_libreoffice(), reason="LibreOffice is not installed")
def test_render_publishes_artifacts_a_manifest_and_a_summary(
    simple_deck: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    output_dir = tmp_path / "out"

    exit_code = main(["render", str(simple_deck), "--output-dir", str(output_dir)])

    assert exit_code == 0
    manifest = json.loads((output_dir / "render.json").read_text(encoding="utf-8"))
    assert manifest["slides"], "the manifest must list rendered slides"
    slide_svg = (output_dir / manifest["slides"][0]["relative_path"]).read_text(encoding="utf-8")
    assert "data:image" not in slide_svg
    assert "rendered" in capsys.readouterr().out


@pytest.mark.skipif(_requires_libreoffice(), reason="LibreOffice is not installed")
def test_render_refuses_a_non_empty_output_directory(simple_deck: Path, tmp_path: Path) -> None:
    output_dir = tmp_path / "out"
    output_dir.mkdir()
    (output_dir / "existing.txt").write_text("keep", encoding="utf-8")

    assert main(["render", str(simple_deck), "--output-dir", str(output_dir)]) == 2
    assert (output_dir / "existing.txt").read_text(encoding="utf-8") == "keep"


def test_render_refuses_a_deck_whose_render_does_not_mirror_its_shapes(
    simple_deck: Path,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    """A converter that drops a shape must fail the CLI, not publish a half-verified deck."""
    output_dir = tmp_path / "out"
    stub = tmp_path / "stub-soffice"
    stub.write_text("", encoding="utf-8")
    truncated = (
        '<svg xmlns="http://www.w3.org/2000/svg" width="33866" height="19050"'
        ' viewBox="0 0 33866 19050"><g class="Slide" id="s1"><g class="Page">'
        '<g class="com.sun.star.drawing.CustomShape"><g id="only">'
        '<rect class="BoundingBox" x="0" y="0" width="10" height="10"/></g></g>'
        "</g></g></svg>"
    )

    class _DroppingConverter:
        def convert(self, source: Path, output_dir: Path) -> Path:
            produced = output_dir / f"{source.stem}.svg"
            produced.write_text(truncated, encoding="utf-8")
            return produced

    monkeypatch.setattr("impromptu_ingestion.cli.discover_soffice", lambda: stub)
    monkeypatch.setattr("impromptu_ingestion.cli.converter_version", lambda _: "stub")
    monkeypatch.setattr(
        "impromptu_ingestion.cli.LibreOfficeSvgConverter", lambda _soffice: _DroppingConverter()
    )

    exit_code = main(["render", str(simple_deck), "--output-dir", str(output_dir)])

    assert exit_code == 2
    assert "slide_mapping_mismatch" in capsys.readouterr().err
    assert list(output_dir.rglob("*")) == [], "a refused render must publish nothing"

    allowed = tmp_path / "allowed"
    assert (
        main(
            [
                "render",
                str(simple_deck),
                "--output-dir",
                str(allowed),
                "--allow-mapping-mismatch",
            ]
        )
        == 0
    )
    assert (allowed / "render.json").is_file()
    assert "animation withheld" in capsys.readouterr().out


def test_doctor_json_reports_renderer_availability(capsys: pytest.CaptureFixture[str]) -> None:
    assert main(["doctor", "--json"]) in (0, 1)

    report = json.loads(capsys.readouterr().out)
    rendering = report["rendering"]
    assert rendering["status"] in ("configured", "not_configured")
    assert isinstance(rendering["renderer"], str | None)
