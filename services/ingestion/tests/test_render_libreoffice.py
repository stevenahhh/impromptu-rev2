from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from impromptu_ingestion.render.libreoffice import (
    LibreOfficeConversionError,
    LibreOfficeSvgConverter,
    SvgConverter,
    converter_version,
    discover_soffice,
    split_slides,
)


def test_converter_protocol_accepts_an_in_memory_fake(tmp_path: Path) -> None:
    class FakeConverter:
        def convert(self, source: Path, output_dir: Path) -> Path:
            del source
            target = output_dir / "deck.svg"
            target.write_text("<svg/>", encoding="utf-8")
            return target

    converter: SvgConverter = FakeConverter()
    assert converter.convert(tmp_path / "deck.pptx", tmp_path).read_text() == "<svg/>"


def test_split_slides_discards_metadata_only_pages_and_preserves_ids() -> None:
    source = """<svg xmlns="http://www.w3.org/2000/svg"
      xmlns:xlink="http://www.w3.org/1999/xlink" width="10" height="20" viewBox="0 0 10 20">
      <defs><clipPath id="clip-original"><rect width="10" height="20"/></clipPath></defs>
      <g class="Slide" id="slide-empty"><g class="Page" id="page-empty">
        <g class="Background"/><g class="Footer"/>
      </g></g>
      <g class="Slide" id="slide-real"><g class="Page" id="page-real">
        <g class="Background"/><g class="TextShape" id="shape-original"/>
      </g></g>
    </svg>"""

    slides = split_slides(source)

    assert len(slides) == 1
    assert 'id="slide-real"' in slides[0]
    assert 'id="shape-original"' in slides[0]
    assert 'id="clip-original"' in slides[0]
    assert "slide-empty" not in slides[0]
    assert 'viewBox="0 0 10 20"' in slides[0]
    assert 'xmlns="http://www.w3.org/2000/svg"' in slides[0]
    assert 'xmlns:xlink="http://www.w3.org/1999/xlink"' in slides[0]


def test_split_slides_rejects_malformed_or_oversized_xml() -> None:
    with pytest.raises(LibreOfficeConversionError, match="invalid_svg") as malformed:
        split_slides("<svg><g>")
    assert malformed.value.code == "invalid_svg"

    with pytest.raises(LibreOfficeConversionError, match="svg_too_large"):
        split_slides(" " * 16_000_001)


def test_converter_reports_missing_output(tmp_path: Path) -> None:
    executable = tmp_path / "soffice"
    executable.write_text("placeholder", encoding="utf-8")
    converter = LibreOfficeSvgConverter(executable, timeout_seconds=5)

    with pytest.MonkeyPatch.context() as monkeypatch:
        monkeypatch.setattr("subprocess.run", lambda *args, **kwargs: _successful_process())
        with pytest.raises(LibreOfficeConversionError) as raised:
            converter.convert(tmp_path / "deck.pptx", tmp_path / "out")

    assert raised.value.code == "libreoffice_output_missing"


def _successful_process() -> object:
    class Result:
        returncode = 0
        stdout = ""
        stderr = ""

    return Result()


def test_discover_soffice_prefers_configured_existing_path(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    executable = tmp_path / "soffice.exe"
    executable.write_bytes(b"")
    monkeypatch.setenv("SOFFICE_PATH", str(executable))

    assert discover_soffice() == executable


def test_converter_version_returns_unknown_for_an_invalid_executable(tmp_path: Path) -> None:
    assert converter_version(tmp_path / "missing-soffice") == "unknown"


_SOFFICE = discover_soffice()


@pytest.mark.skipif(_SOFFICE is None, reason="LibreOffice is not installed")
def test_real_libreoffice_converts_and_splits_a_presentation(tmp_path: Path) -> None:
    assert _SOFFICE is not None
    source = tmp_path / "integration.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "Rendered"
    deck.save(source)

    output = LibreOfficeSvgConverter(_SOFFICE, timeout_seconds=60).convert(source, tmp_path / "out")
    slides = split_slides(output.read_text(encoding="utf-8"))

    assert output.is_file()
    assert len(slides) == 1
    assert "Rendered" in slides[0]
    assert converter_version(_SOFFICE).startswith("LibreOffice ")
