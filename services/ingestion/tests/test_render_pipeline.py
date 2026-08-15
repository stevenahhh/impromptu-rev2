import base64
import json
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from impromptu_ingestion.render.contracts import RenderedDeck
from impromptu_ingestion.render.libreoffice import SvgConverter
from impromptu_ingestion.render.pipeline import RenderRequest, render_deck

_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def _two_shape_deck(path: Path) -> tuple[int, int]:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    first = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(1), Inches(3), Inches(2)
    )
    first.name = "카드A"
    first.text_frame.text = "한국어"
    second = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(6), Inches(4), Inches(3), Inches(2))
    second.name = "카드B"
    presentation.save(path)
    return first.shape_id, second.shape_id


def _svg_for(
    shape_ids: tuple[int, int], *, drop_second: bool = False, strip_ids: bool = False
) -> str:
    containers = [
        f"""<g class="com.sun.star.drawing.CustomShape" id="idcontainer{index}">
             <g id="id{index + 3}">
               <rect class="BoundingBox" x="{x}" y="{y}" width="{width}" height="{height}"/>
               <image xlink:href="data:image/png;base64,{base64.b64encode(_PNG).decode()}"
                      x="{x}" y="{y}" width="8" height="8"/>
             </g>
           </g>"""
        for index, (x, y, width, height) in enumerate(
            ((2540, 2540, 7620, 5080), (15240, 10160, 7620, 5080))
        )
    ]
    if drop_second:
        containers = containers[:1]
    if strip_ids:
        containers = [
            container.replace(f'<g id="id{index + 3}">', "<g>").replace(
                f' id="idcontainer{index}"', ""
            )
            for index, container in enumerate(containers)
        ]
    return (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<svg xmlns="http://www.w3.org/2000/svg" xmlns:xlink="http://www.w3.org/1999/xlink"'
        ' width="33866" height="19050" viewBox="0 0 33866 19050">'
        '<g class="Slide" id="slide1"><g class="Page">' + "".join(containers) + "</g></g></svg>"
    )


class _FakeConverter(SvgConverter):
    def __init__(self, svg: str) -> None:
        self._svg = svg

    def convert(self, source: Path, output_dir: Path) -> Path:
        produced = output_dir / f"{source.stem}.svg"
        produced.write_text(self._svg, encoding="utf-8")
        return produced


def _request(source: Path, output_dir: Path, converter: SvgConverter) -> RenderRequest:
    return RenderRequest(source=source, output_dir=output_dir, converter=converter)


def test_render_publishes_externalized_artifacts_and_a_manifest(tmp_path: Path) -> None:
    deck_path = tmp_path / "deck.pptx"
    shape_ids = _two_shape_deck(deck_path)
    output_dir = tmp_path / "out"

    rendered = render_deck(_request(deck_path, output_dir, _FakeConverter(_svg_for(shape_ids))))

    assert isinstance(rendered, RenderedDeck)
    assert len(rendered.slides) == 1
    slide_svg = (output_dir / rendered.slides[0].relative_path).read_text(encoding="utf-8")
    assert "data:image" not in slide_svg
    assert rendered.assets, "the inline image must be externalized to a file"
    for asset in rendered.assets:
        assert (output_dir / asset.relative_path).is_file()
    manifest = json.loads((output_dir / "render.json").read_text(encoding="utf-8"))
    assert manifest["deck_id"] == rendered.deck_id


def test_structural_mismatch_publishes_nothing_and_fails_with_a_typed_code(tmp_path: Path) -> None:
    """A deck whose render does not mirror its shapes must not reach the output directory."""
    deck_path = tmp_path / "deck.pptx"
    shape_ids = _two_shape_deck(deck_path)
    output_dir = tmp_path / "out"

    with pytest.raises(Exception) as failure:
        render_deck(
            _request(deck_path, output_dir, _FakeConverter(_svg_for(shape_ids, drop_second=True)))
        )

    assert getattr(failure.value, "code", "") == "slide_mapping_mismatch"
    assert "top_level_count_mismatch" in str(failure.value)
    assert list(output_dir.rglob("*")) == [], "a refused render must leave no artifacts behind"


def test_mismatch_may_be_published_static_only_when_explicitly_allowed(tmp_path: Path) -> None:
    deck_path = tmp_path / "deck.pptx"
    shape_ids = _two_shape_deck(deck_path)
    output_dir = tmp_path / "out"

    rendered = render_deck(
        RenderRequest(
            source=deck_path,
            output_dir=output_dir,
            converter=_FakeConverter(_svg_for(shape_ids, drop_second=True)),
            strict_mapping=False,
        )
    )

    assert rendered.animation_eligible is False
    assert rendered.ineligible_reason is not None
    assert rendered.mapping_issues, "a dropped container must be reported as a mapping issue"
    assert (output_dir / rendered.slides[0].relative_path).is_file()


def test_containers_without_a_renderer_id_are_stamped_and_stay_targetable(tmp_path: Path) -> None:
    """LibreOffice emits no id for Graphic/TableShape containers; we must supply our own."""
    deck_path = tmp_path / "deck.pptx"
    shape_ids = _two_shape_deck(deck_path)
    output_dir = tmp_path / "out"

    rendered = render_deck(
        _request(deck_path, output_dir, _FakeConverter(_svg_for(shape_ids, strip_ids=True)))
    )

    assert [issue.code for issue in rendered.mapping_issues] == []
    written = (output_dir / rendered.slides[0].relative_path).read_text(encoding="utf-8")
    assert 'id="impromptu-s1-1"' in written
    assert 'id="impromptu-s1-2"' in written


def test_libreoffice_doctype_prologue_is_normalized_before_parsing(tmp_path: Path) -> None:
    """LibreOffice always emits an XML declaration and an SVG DOCTYPE; both must be tolerated."""
    deck_path = tmp_path / "deck.pptx"
    shape_ids = _two_shape_deck(deck_path)
    prologue = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<!DOCTYPE svg PUBLIC "-//W3C//DTD SVG 1.1//EN"'
        ' "http://www.w3.org/Graphics/SVG/1.1/DTD/svg11.dtd">\n'
    )
    svg = prologue + _svg_for(shape_ids).split("?>", 1)[1]

    rendered = render_deck(_request(deck_path, tmp_path / "out", _FakeConverter(svg)))

    assert len(rendered.slides) == 1


def test_render_refuses_to_overwrite_an_existing_output_directory(tmp_path: Path) -> None:
    deck_path = tmp_path / "deck.pptx"
    shape_ids = _two_shape_deck(deck_path)
    output_dir = tmp_path / "out"
    output_dir.mkdir()
    (output_dir / "keep.txt").write_text("existing", encoding="utf-8")

    with pytest.raises(Exception) as failure:
        render_deck(_request(deck_path, output_dir, _FakeConverter(_svg_for(shape_ids))))

    assert getattr(failure.value, "code", "") == "output_directory_not_empty"
    assert (output_dir / "keep.txt").read_text(encoding="utf-8") == "existing"
