import hashlib
import json
import shutil
import subprocess
from contextlib import AbstractContextManager
from pathlib import Path
from typing import Any, cast

import pymupdf
import pytest

from impromptu_ingestion import vision
from impromptu_ingestion.adapters import (
    PdfStructuralAdapter,
    PptxStructuralAdapter,
    RenderingUnsupportedError,
    StructuralExtractionError,
)
from impromptu_ingestion.adapters.pdf import extract_pdf_in_process
from impromptu_ingestion.canonical import canonical_manifest_bytes, manifest_sha256
from impromptu_ingestion.contracts import (
    DeckManifest,
    IngestionJob,
    IngestionLimits,
    InputKind,
    TextElement,
    ValidatedInput,
)
from impromptu_ingestion.validation import stage_input

_REPOSITORY_ROOT = Path(__file__).resolve().parents[3]
_FROZEN_FIXTURE_ROOT = _REPOSITORY_ROOT / "tests" / "fixtures" / "format-neutral-decks"
_FROZEN_FIXTURE_REGISTRY = _FROZEN_FIXTURE_ROOT / "fixture-manifest.json"


def _fixture_registry() -> dict[str, Any]:
    return json.loads(_FROZEN_FIXTURE_REGISTRY.read_text(encoding="utf-8"))


def _validate_fixture_hashes(registry: dict[str, Any], repository_root: Path) -> None:
    for fixture in registry["fixtures"]:
        path = repository_root / fixture["path"]
        actual = hashlib.sha256(path.read_bytes()).hexdigest()
        assert actual == fixture["sha256"], f"SHA-256 mismatch for {fixture['path']}"


def _staged(
    path: Path, limits: IngestionLimits | None = None
) -> AbstractContextManager[ValidatedInput]:
    return stage_input(
        IngestionJob(
            job_id="job_adapter_01",
            source=path,
            limits=limits or IngestionLimits(),
        )
    )


def _write_raster_only_pdf(path: Path) -> None:
    """One full-page image, no text layer: the canonical scanned-page fixture."""

    document = pymupdf.open()
    page = document.new_page(width=720, height=405)
    raster = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 16, 16))
    page.insert_image(pymupdf.Rect(0, 0, 720, 405), stream=raster.tobytes("png"))
    document.save(path)


def test_pptx_extracts_text_table_image_and_chart_without_private_notes(
    sample_pptx: Path,
) -> None:
    with _staged(sample_pptx) as source:
        manifest = PptxStructuralAdapter().extract(source)

    assert manifest.source_kind is InputKind.PPTX
    assert len(manifest.slides) == 1
    assert {element.kind for element in manifest.slides[0].elements} == {
        "text",
        "table",
        "image",
        "chart",
    }
    assert b"PRIVATE-NOTES-MUST-NEVER-LEAK" not in canonical_manifest_bytes(manifest)
    assert manifest.render_boundary.fidelity_verified is False


def test_adapter_parses_the_hashed_stage_after_source_replacement(sample_pptx: Path) -> None:
    original_digest = hashlib.sha256(sample_pptx.read_bytes()).hexdigest()

    with _staged(sample_pptx) as source:
        sample_pptx.write_bytes(b"replacement after staging")
        manifest = PptxStructuralAdapter().extract(source)

    assert manifest.source_sha256 == original_digest
    assert any(
        element.kind == "text" and element.text == "한국어 근거 자료"
        for element in manifest.slides[0].elements
    )


def test_pptx_manifest_and_identity_are_repeatable(sample_pptx: Path) -> None:
    adapter = PptxStructuralAdapter()
    with _staged(sample_pptx) as source:
        first = adapter.extract(source)
        second = adapter.extract(source)

    assert canonical_manifest_bytes(first) == canonical_manifest_bytes(second)
    assert manifest_sha256(first) == manifest_sha256(second)
    assert first.deck_id.startswith("deck_")
    assert first.slides[0].slide_key.startswith("slide_")


def test_pdf_extracts_positioned_text_and_images(sample_pdf: Path) -> None:
    with _staged(sample_pdf) as source:
        manifest = PdfStructuralAdapter().extract(source)

    assert manifest.source_kind is InputKind.PDF
    assert len(manifest.slides) == 2
    assert any(element.kind == "text" for element in manifest.slides[0].elements)
    assert any(element.kind == "image" for element in manifest.slides[0].elements)
    assert any(element.kind == "text" for element in manifest.slides[1].elements)
    assert manifest.slides[1].warnings == ()


def test_frozen_pdf_and_pptx_share_korean_structural_content() -> None:
    registry = _fixture_registry()
    sentinel = registry["sentinels"]["structural"]
    forbidden = registry["sentinels"]["forbiddenPptxText"]
    density = registry["density"]
    fixture_by_kind = {fixture["kind"]: fixture for fixture in registry["fixtures"]}

    manifests: dict[str, DeckManifest] = {}
    texts_by_kind: dict[str, tuple[str, ...]] = {}
    for kind, adapter in (
        ("pptx", PptxStructuralAdapter()),
        ("text-layer-pdf", PdfStructuralAdapter()),
    ):
        path = _REPOSITORY_ROOT / fixture_by_kind[kind]["path"]
        with _staged(path) as source:
            manifest = adapter.extract(source)
        manifests[kind] = manifest
        texts_by_kind[kind] = tuple(
            element.text
            for slide in manifest.slides
            for element in slide.elements
            if element.kind == "text"
        )

    # Format neutrality: both formats must yield byte-identical ordered text.
    assert texts_by_kind["pptx"] == texts_by_kind["text-layer-pdf"]
    assert texts_by_kind["pptx"][0] == sentinel
    assert forbidden not in "".join(texts_by_kind["pptx"])

    # Pinned density invariants from the frozen registry.
    assert len(manifests["pptx"].slides) == density["slideCount"]
    assert len(manifests["text-layer-pdf"].slides) == density["slideCount"]
    assert len(texts_by_kind["pptx"]) == density["textElementCount"]
    assert sum(len(text) for text in texts_by_kind["pptx"]) == density["totalTextCharacters"]


def test_frozen_scanned_pdf_uses_ocr_only_for_the_raster_page(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    registry = _fixture_registry()
    fixture = next(item for item in registry["fixtures"] if item["kind"] == "scanned-pdf")
    sentinel = registry["sentinels"]["ocr"]
    calls: list[int] = []

    def extracted(page: object, *, timeout_seconds: float) -> tuple[TextElement, ...]:
        calls.append(round(timeout_seconds))
        return (
            TextElement(
                element_id="ocr:text:1",
                text=sentinel,
                x=72,
                y=72,
                width=300,
                height=30,
            ),
        )

    monkeypatch.setattr("impromptu_ingestion.adapters.pdf.extract_ocr_text", extracted)
    with _staged(_REPOSITORY_ROOT / fixture["path"]) as source:
        manifest = extract_pdf_in_process(source)

    assert calls == [30] * registry["density"]["slideCount"]
    assert [
        [element.text for element in slide.elements if element.kind == "text"]
        for slide in manifest.slides
    ] == [[sentinel]] * registry["density"]["slideCount"]
    assert {warning.code for slide in manifest.slides for warning in slide.warnings} == {
        "ocr_applied"
    }


def test_pdf_prefers_embedded_text_layer_over_ocr(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A page carrying real text is never routed through OCR, image or not."""

    def fail(page: object, *, timeout_seconds: float) -> tuple[TextElement, ...]:
        raise AssertionError("OCR must not run on a page with an embedded text layer")

    monkeypatch.setattr("impromptu_ingestion.adapters.pdf.extract_ocr_text", fail)
    mixed = tmp_path / "mixed.pdf"
    document = pymupdf.open()
    page = document.new_page(width=720, height=405)
    raster = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 16, 16))
    page.insert_image(pymupdf.Rect(0, 0, 720, 405), stream=raster.tobytes("png"))
    page.insert_text((72, 72), "임베디드 텍스트 2026", fontname="korea")
    document.save(mixed)

    with _staged(mixed) as source:
        manifest = PdfStructuralAdapter().extract(source)

    texts = [element.text for element in manifest.slides[0].elements if element.kind == "text"]
    assert any("임베디드 텍스트 2026" in text for text in texts)
    assert manifest.slides[0].warnings == ()


def test_raster_page_prefers_vision_transcription_when_configured(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.delenv("TESSDATA_PREFIX", raising=False)
    scanned = tmp_path / "scanned.pdf"
    _write_raster_only_pdf(scanned)

    def fail(page: object, *, timeout_seconds: float) -> tuple[TextElement, ...]:
        raise AssertionError("tesseract must not run when vision transcription succeeded")

    monkeypatch.setattr("impromptu_ingestion.adapters.pdf.extract_ocr_text", fail)
    monkeypatch.setattr(
        "impromptu_ingestion.adapters.pdf.vision_config",
        lambda: vision.VisionConfig(
            base_url="https://vision.example/v1",
            api_key="secret",
            model="vision-model-x",
        ),
    )

    transcribe_calls: list[object] = []

    def fake_transcribe(jobs, **options):
        transcribe_calls.append({"jobs": jobs, "options": options})
        return {
            index: vision._line_elements(
                [f"비전 텍스트 {index}"], page_width=720.0, page_height=405.0
            )
            for index in jobs
        }

    monkeypatch.setattr("impromptu_ingestion.adapters.pdf.transcribe_pages", fake_transcribe)

    with _staged(scanned) as source:
        manifest = extract_pdf_in_process(source)

    assert len(transcribe_calls) == 1
    assert sorted(transcribe_calls[0]["jobs"].keys()) == [0]
    texts = [element.text for element in manifest.slides[0].elements if element.kind == "text"]
    assert texts == ["비전 텍스트 0"]
    warning_codes = {warning.code for warning in manifest.slides[0].warnings}
    assert warning_codes == {"vision_ocr_applied"}


def test_raster_page_falls_back_to_tesseract_when_vision_fails(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("TESSDATA_PREFIX", "/opt/homebrew/share/tessdata")
    scanned = tmp_path / "scanned.pdf"
    _write_raster_only_pdf(scanned)

    monkeypatch.setattr(
        "impromptu_ingestion.adapters.pdf.vision_config",
        lambda: vision.VisionConfig(
            base_url="https://vision.example/v1",
            api_key="secret",
            model="vision-model-x",
        ),
    )
    monkeypatch.setattr(
        "impromptu_ingestion.adapters.pdf.transcribe_pages", lambda jobs, **options: {}
    )

    def extracted(page: object, *, timeout_seconds: float) -> tuple[TextElement, ...]:
        return (
            TextElement(
                element_id="ocr:text:1",
                text="테서랙트 텍스트",
                x=1,
                y=1,
                width=10,
                height=10,
            ),
        )

    monkeypatch.setattr("impromptu_ingestion.adapters.pdf.extract_ocr_text", extracted)

    with _staged(scanned) as source:
        manifest = extract_pdf_in_process(source)

    texts = [element.text for element in manifest.slides[0].elements if element.kind == "text"]
    assert texts == ["테서랙트 텍스트"]
    warning_codes = {warning.code for warning in manifest.slides[0].warnings}
    assert warning_codes == {"ocr_applied"}


def test_scanned_page_ingests_with_warning_when_ocr_is_unavailable(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.delenv("TESSDATA_PREFIX", raising=False)
    scanned = tmp_path / "scanned.pdf"
    _write_raster_only_pdf(scanned)

    with _staged(scanned) as source:
        manifest = PdfStructuralAdapter().extract(source)

    assert len(manifest.slides) == 1
    assert any(element.kind == "image" for element in manifest.slides[0].elements)
    assert not any(element.kind == "text" for element in manifest.slides[0].elements)
    warning_codes = {warning.code for warning in manifest.slides[0].warnings}
    assert "ocr_unavailable" in warning_codes


def test_fixture_hash_registry_accepts_frozen_binaries() -> None:
    _validate_fixture_hashes(_fixture_registry(), _REPOSITORY_ROOT)


def test_fixture_hash_registry_rejects_mutated_binary(tmp_path: Path) -> None:
    registry = _fixture_registry()
    copied_root = tmp_path / "repository"
    copied_fixtures = copied_root / _FROZEN_FIXTURE_ROOT.relative_to(_REPOSITORY_ROOT)
    shutil.copytree(_FROZEN_FIXTURE_ROOT, copied_fixtures)
    mutated = copied_root / registry["fixtures"][0]["path"]
    mutated.write_bytes(mutated.read_bytes() + b"fixture mutation")

    with pytest.raises(AssertionError, match="SHA-256 mismatch"):
        _validate_fixture_hashes(registry, copied_root)


def test_fixture_registry_reserves_only_private_0008_and_0009() -> None:
    reservations = _fixture_registry()["migrationReservations"]
    assert reservations == {
        "private": [
            {"filename": "0008_deck_retrieval_hybrid.sql", "owner": "retrieval"},
            {"filename": "0009_session_reports.sql", "owner": "report"},
        ],
        "projection": {"latest": "0005_gateway_state.sql", "newMigrationAllowed": False},
    }


@pytest.mark.parametrize("adapter", [PptxStructuralAdapter(), PdfStructuralAdapter()])
def test_rendering_is_an_explicit_unsupported_boundary(
    adapter: PptxStructuralAdapter | PdfStructuralAdapter,
) -> None:
    with pytest.raises(RenderingUnsupportedError, match="not configured"):
        adapter.render_slides()


def test_truncated_pdf_that_mupdf_can_repair_is_strictly_rejected(
    sample_pdf: Path, tmp_path: Path
) -> None:
    truncated = tmp_path / "truncated.pdf"
    original = sample_pdf.read_bytes()
    assert original.rstrip().endswith(b"%%EOF")
    truncated.write_bytes(original[: original.rfind(b"%%EOF")])

    with _staged(truncated) as staged, pytest.raises(StructuralExtractionError) as raised:
        PdfStructuralAdapter().extract(staged)

    assert raised.value.code == "repair_required"


def test_pdf_enforces_page_object_and_resource_limits(sample_pdf: Path) -> None:
    cases = (
        IngestionLimits(max_pdf_pages=1),
        IngestionLimits(max_pdf_objects=1),
        IngestionLimits(max_pdf_resource_bytes=1),
        IngestionLimits(max_pdf_elements=1),
        IngestionLimits(max_pdf_image_pixels=1),
    )

    for limits in cases:
        with (
            _staged(sample_pdf, limits) as staged,
            pytest.raises(StructuralExtractionError) as raised,
        ):
            PdfStructuralAdapter().extract(staged)
        assert raised.value.code == "resource_limit"


def test_pdf_worker_timeout_is_a_typed_failure(
    sample_pdf: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def timeout(*args: object, **kwargs: object) -> subprocess.CompletedProcess[bytes]:
        raise subprocess.TimeoutExpired(cmd="pdf-worker", timeout=1)

    monkeypatch.setattr(subprocess, "run", timeout)

    with (
        _staged(sample_pdf, IngestionLimits(operation_timeout_seconds=1)) as staged,
        pytest.raises(StructuralExtractionError) as raised,
    ):
        PdfStructuralAdapter().extract(staged)

    assert raised.value.code == "operation_timeout"


def test_malformed_pdf_fails_with_a_typed_extraction_error(tmp_path: Path) -> None:
    source = tmp_path / "malformed.pdf"
    source.write_bytes(b"%PDF-1.7\nthis is not a document")

    with _staged(source) as staged, pytest.raises(StructuralExtractionError) as raised:
        PdfStructuralAdapter().extract(staged)

    assert raised.value.code == "invalid_document"


def test_a_picture_linked_to_an_external_file_is_skipped_not_fatal() -> None:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    from impromptu_ingestion.adapters.pptx import _extract_shape

    class _LinkedPicture:
        shape_id = 7
        left = 914400
        top = 914400
        width = 914400
        height = 914400
        has_table = False
        has_chart = False
        has_text_frame = False
        shape_type = MSO_SHAPE_TYPE.PICTURE

        @property
        def image(self) -> object:
            # What python-pptx raises for a picture whose bytes live outside the package.
            raise ValueError("no embedded image")

    # There are no bytes to hash, so the shape cannot become an element — but it must be skipped
    # and warned about, not raised. One linked picture on one slide used to reduce an entire
    # 21-slide deck to `invalid_document` and fail the upload outright.
    assert _extract_shape(cast(Any, _LinkedPicture())) is None
