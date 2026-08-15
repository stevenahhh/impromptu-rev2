from collections.abc import Iterable

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.shapes.base import BaseShape
from pptx.util import Inches

from impromptu_ingestion.render.mapping import map_slide


def _xml(slide: object) -> bytes:
    return etree.tostring(slide.element)  # type: ignore[attr-defined, no-any-return]


def _rect(shape: BaseShape) -> str:
    return (
        f'<rect class="BoundingBox" x="{shape.left / 360}" y="{shape.top / 360}" '
        f'width="{shape.width / 360}" height="{shape.height / 360}"/>'
    )


def _svg(containers: Iterable[str]) -> str:
    return (
        '<svg xmlns="http://www.w3.org/2000/svg"><g class="Slide"><g class="Page">'
        + "".join(containers)
        + "</g></g></svg>"
    )


def _container(kind: str, element_id: str, shape: BaseShape, children: str = "") -> str:
    return f'<g class="{kind}"><g id="{element_id}">{_rect(shape)}</g>{children}</g>'


def test_clean_nested_group_resolves_every_target() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    outer = slide.shapes.add_group_shape()
    first = outer.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
    inner = outer.shapes.add_group_shape()
    second = inner.shapes.add_shape(MSO_SHAPE.OVAL, Inches(4), Inches(2), Inches(1), Inches(1))
    inner_children = _container("com.sun.star.drawing.CustomShape", "id3", second)
    outer_children = _container("com.sun.star.drawing.CustomShape", "id1", first) + _container(
        "Group", "id2", inner, inner_children
    )

    result = map_slide(_xml(slide), _svg([_container("Group", "id0", outer, outer_children)]), 1)

    assert result.issues == ()
    assert [(target.shape_id, target.svg_element_id) for target in result.targets] == [
        (outer.shape_id, "id0"),
        (first.shape_id, "id1"),
        (inner.shape_id, "id2"),
        (second.shape_id, "id3"),
    ]


def test_hidden_shape_is_skipped_without_desynchronizing() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    hidden = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(1), Inches(1))
    hidden._element.nvSpPr.cNvPr.set("hidden", "1")  # type: ignore[attr-defined]
    visible = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(3), Inches(1), Inches(1), Inches(1))

    result = map_slide(
        _xml(slide),
        _svg([_container("com.sun.star.drawing.CustomShape", "id1", visible)]),
        1,
    )

    assert result.issues == ()
    assert [(target.shape_id, target.svg_element_id) for target in result.targets] == [
        (visible.shape_id, "id1")
    ]


def test_empty_placeholder_is_kept() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title, subtitle = slide.shapes

    result = map_slide(
        _xml(slide),
        _svg(
            [
                _container("TitleText", "id1", title),
                _container("Outline", "id2", subtitle),
            ]
        ),
        1,
    )

    assert result.issues == ()
    assert result.targets[0].shape_id == title.shape_id


def test_title_and_body_placeholders_use_global_placeholder_order() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    title, body = slide.shapes

    result = map_slide(
        _xml(slide),
        _svg(
            [
                _container("TitleText", "id1", title),
                _container("Outline", "id2", body),
            ]
        ),
        2,
    )

    assert result.issues == ()
    assert [target.svg_element_id for target in result.targets] == ["id1", "id2"]


def test_missing_group_child_records_child_count_mismatch() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    child = group.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(1), Inches(1))

    result = map_slide(_xml(slide), _svg([_container("Group", "id1", group)]), 1)

    assert [(issue.code, issue.path) for issue in result.issues] == [
        ("child_count_mismatch", "slide1.0")
    ]
    assert child.shape_id not in {target.shape_id for target in result.targets}


def test_shifted_container_records_center_mismatch() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(1), Inches(1))
    shifted = (
        '<g class="com.sun.star.drawing.CustomShape"><g id="id1">'
        '<rect class="BoundingBox" x="5000" y="2540" width="2540" height="2540"/>'
        "</g></g>"
    )

    result = map_slide(_xml(slide), _svg([shifted]), 3)

    assert [(issue.code, issue.path) for issue in result.issues] == [
        ("center_mismatch", "slide3.0")
    ]


def test_wrong_container_type_records_type_mismatch() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(1), Inches(1))

    result = map_slide(_xml(slide), _svg([_container("Graphic", "id1", shape)]), 4)

    assert [(issue.code, issue.path) for issue in result.issues] == [("type_mismatch", "slide4.0")]
