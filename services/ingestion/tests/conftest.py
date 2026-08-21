import base64
from io import BytesIO
from pathlib import Path

import pymupdf
import pytest
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


@pytest.fixture
def sample_pptx(tmp_path: Path) -> Path:
    path = tmp_path / "한국어-구조.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])

    text_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(4), Inches(0.5))
    text_box.text_frame.text = "한국어 근거 자료"

    table_shape = slide.shapes.add_table(2, 2, Inches(0.5), Inches(1.2), Inches(4), Inches(1))
    table_shape.table.cell(0, 0).text = "연도"
    table_shape.table.cell(0, 1).text = "값"
    table_shape.table.cell(1, 0).text = "2026"
    table_shape.table.cell(1, 1).text = "42"

    slide.shapes.add_picture(BytesIO(_PNG), Inches(5), Inches(0.5), Inches(1), Inches(1))

    chart_data = ChartData()
    chart_data.categories = ("상반기", "하반기")
    chart_data.add_series("매출", (10, 20))
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.5),
        Inches(2.5),
        Inches(5),
        Inches(2.5),
        chart_data,
    )
    slide.notes_slide.notes_text_frame.text = "PRIVATE-NOTES-MUST-NEVER-LEAK"
    presentation.save(path)
    return path


@pytest.fixture
def sample_pdf(tmp_path: Path) -> Path:
    path = tmp_path / "구조.pdf"
    document = pymupdf.open()
    first = document.new_page(width=720, height=405)
    first.insert_text((72, 72), "Evidence 2026")
    first.insert_image(pymupdf.Rect(100, 100, 110, 110), stream=_PNG)

    second = document.new_page(width=720, height=405)
    second.insert_text((72, 72), "Structural fallback")
    second.insert_image(pymupdf.Rect(0, 0, 10, 10), stream=_PNG)
    document.save(path)
    document.close()
    return path
