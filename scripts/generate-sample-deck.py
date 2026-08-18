# /// script
# requires-python = ">=3.14"
# dependencies = ["python-pptx==1.0.2"]
# ///
# ─── How to run ───
# uv run scripts/generate-sample-deck.py

from __future__ import annotations

import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
from typing import Final
from xml.etree import ElementTree

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.presentation import Presentation as PresentationType
from pptx.slide import Slide
from pptx.util import Inches, Pt

ROOT: Final = Path(__file__).resolve().parents[1]
OUTPUT_DIR: Final = ROOT / "docs" / "samples"
PPTX_PATH: Final = OUTPUT_DIR / "impromptu-sample-deck.pptx"
PDF_PATH: Final = OUTPUT_DIR / "impromptu-sample-deck.pdf"

INK: Final = RGBColor(24, 24, 27)
MUTED: Final = RGBColor(113, 113, 122)
CANVAS: Final = RGBColor(250, 250, 250)
CARD: Final = RGBColor(255, 255, 255)
BORDER: Final = RGBColor(228, 228, 231)
PRIMARY: Final = RGBColor(24, 24, 27)
PRIMARY_INK: Final = RGBColor(250, 250, 250)
ACCENT: Final = RGBColor(37, 99, 235)
SUCCESS: Final = RGBColor(22, 163, 74)
FONT: Final = "Apple SD Gothic Neo"
MONO: Final = "Menlo"


def add_text(
    slide: Slide,
    text: str,
    left: float,
    top: float,
    width: float,
    height: float,
    *,
    size: int,
    color: RGBColor = INK,
    bold: bool = False,
    font: str = FONT,
    align: PP_ALIGN = PP_ALIGN.LEFT,
) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_card(
    slide: Slide,
    left: float,
    top: float,
    width: float,
    height: float,
    *,
    fill: RGBColor = CARD,
    border: RGBColor = BORDER,
    radius: MSO_SHAPE = MSO_SHAPE.ROUNDED_RECTANGLE,
) -> None:
    shape = slide.shapes.add_shape(
        radius,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = border
    shape.line.width = Pt(1)


def add_header(slide: Slide, section: str, page: int) -> None:
    add_text(slide, "IMPROMPTU", 0.7, 0.25, 2.2, 0.35, size=10, bold=True)
    add_text(slide, section, 9.3, 0.25, 2.8, 0.35, size=10, color=MUTED, align=PP_ALIGN.RIGHT)
    add_text(slide, f"{page:02d}", 12.15, 0.25, 0.5, 0.35, size=10, color=MUTED, font=MONO)


def add_title(slide: Slide, eyebrow: str, title: str, body: str) -> None:
    add_text(slide, eyebrow, 0.75, 1.0, 4.0, 0.35, size=11, color=ACCENT, bold=True)
    add_text(slide, title, 0.75, 1.45, 11.8, 1.15, size=31, bold=True)
    add_text(slide, body, 0.78, 2.55, 10.7, 0.75, size=15, color=MUTED)


def base_slide(deck: PresentationType, section: str, page: int) -> Slide:
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = CANVAS
    add_header(slide, section, page)
    return slide


def build_deck() -> PresentationType:
    deck = Presentation()
    deck.slide_width = Inches(13.333)
    deck.slide_height = Inches(7.5)
    deck.core_properties.title = "Impromptu 샘플 발표"
    deck.core_properties.subject = "업로드부터 발표 시작까지의 제품 흐름"
    deck.core_properties.author = "Impromptu"

    slide = base_slide(deck, "OPENING", 1)
    add_text(slide, "발표가 흐르도록.", 0.78, 1.3, 10.5, 1.4, size=42, bold=True)
    add_text(
        slide,
        "파일 하나로 슬라이드, 발표자 콘솔, 근거 준비, 외장 화면을 연결합니다.",
        0.82,
        2.75,
        9.8,
        0.8,
        size=18,
        color=MUTED,
    )
    add_card(slide, 0.82, 4.55, 3.45, 1.25, fill=PRIMARY, border=PRIMARY)
    add_text(slide, "PPTX / PDF 업로드", 1.15, 4.88, 2.8, 0.55, size=15, color=PRIMARY_INK, bold=True)
    add_text(slide, "DEMO DECK · 2026", 0.82, 6.55, 3.2, 0.35, size=10, color=MUTED, font=MONO)

    slide = base_slide(deck, "PROBLEM", 2)
    add_title(
        slide,
        "발표 전 10분",
        "도구를 준비하느라\n발표를 놓칩니다.",
        "파일, 발표자 노트, 근거 검색, 화면 연결이 서로 다른 도구에 흩어져 있습니다.",
    )
    for index, (number, label) in enumerate(
        (("04", "도구 전환"), ("07", "설정 단계"), ("00", "하나의 기준 화면"))
    ):
        left = 0.8 + index * 4.15
        add_card(slide, left, 4.2, 3.65, 1.6)
        add_text(slide, number, left + 0.3, 4.48, 1.0, 0.55, size=24, bold=True, font=MONO)
        add_text(slide, label, left + 0.3, 5.08, 2.9, 0.4, size=13, color=MUTED)

    slide = base_slide(deck, "FLOW", 3)
    add_title(
        slide,
        "한 화면, 세 단계",
        "올리고. 연결하고. 시작합니다.",
        "준비 상태가 한 방향으로 흐르기 때문에 발표자는 다음 행동만 확인하면 됩니다.",
    )
    for index, (step, title, detail) in enumerate(
        (
            ("01", "슬라이드 준비", "PPTX 또는 PDF를 놓습니다."),
            ("02", "화면 연결", "Stage와 외장 디스플레이를 확인합니다."),
            ("03", "발표 시작", "하나의 버튼으로 발표를 엽니다."),
        )
    ):
        left = 0.8 + index * 4.15
        add_card(slide, left, 4.05, 3.65, 1.85, fill=PRIMARY if index == 0 else CARD)
        color = PRIMARY_INK if index == 0 else INK
        secondary = RGBColor(212, 212, 216) if index == 0 else MUTED
        add_text(slide, step, left + 0.3, 4.27, 0.7, 0.4, size=11, color=secondary, font=MONO)
        add_text(slide, title, left + 0.3, 4.72, 3.0, 0.4, size=16, color=color, bold=True)
        add_text(slide, detail, left + 0.3, 5.17, 3.0, 0.48, size=11, color=secondary)

    slide = base_slide(deck, "EVIDENCE", 4)
    add_title(
        slide,
        "발표는 기다리지 않습니다",
        "AI 근거 준비는\n뒤에서 계속됩니다.",
        "업로드 직후 발표 화면은 열리고, 근거는 상태와 출처를 가진 카드로 준비됩니다.",
    )
    add_card(slide, 7.55, 1.15, 4.85, 4.95)
    add_text(slide, "근거 준비", 7.95, 1.55, 2.6, 0.45, size=17, bold=True)
    rows = (("슬라이드 분석", "완료", SUCCESS), ("출처 교차 확인", "진행 중", ACCENT), ("발표자 승인", "대기", MUTED))
    for index, (label, state, color) in enumerate(rows):
        top = 2.35 + index * 1.05
        add_card(slide, 7.95, top, 4.05, 0.78, fill=CANVAS)
        add_text(slide, label, 8.2, top + 0.16, 2.25, 0.35, size=12)
        add_text(slide, state, 10.45, top + 0.16, 1.2, 0.35, size=11, color=color, bold=True, align=PP_ALIGN.RIGHT)
    add_text(slide, "발표 준비 완료", 7.95, 5.45, 4.05, 0.35, size=12, color=SUCCESS, bold=True)

    slide = base_slide(deck, "DISPLAY", 5)
    add_title(
        slide,
        "청중에게는 Stage만",
        "외장 화면은 공개 정보만\n표시합니다.",
        "발표자 콘솔과 청중 화면을 분리하고, 화면 구성에 맞춰 안전한 모드를 선택합니다.",
    )
    for index, (title, detail) in enumerate(
        (("확장", "발표자와 청중 화면 분리"), ("복제", "공개 화면 안전장치 적용"), ("단일 화면", "Stage 중심 대체 흐름"))
    ):
        top = 3.72 + index * 0.85
        add_card(slide, 7.45, top, 4.9, 0.62, fill=CARD)
        add_text(slide, title, 7.72, top + 0.1, 1.2, 0.35, size=12, bold=True)
        add_text(slide, detail, 8.95, top + 0.1, 3.05, 0.35, size=11, color=MUTED)

    slide = base_slide(deck, "CONTROL", 6)
    add_title(
        slide,
        "발표자 콘솔",
        "슬라이드와 근거를\n같은 자리에서 제어합니다.",
        "현재 슬라이드, 다음 슬라이드, 승인된 근거, 화면 상태가 발표자에게만 보입니다.",
    )
    add_card(slide, 7.15, 1.25, 5.15, 4.75, fill=PRIMARY, border=PRIMARY)
    add_text(slide, "03 / 07", 7.55, 1.62, 1.4, 0.38, size=11, color=RGBColor(161, 161, 170), font=MONO)
    add_text(slide, "한 화면,\n세 단계", 7.55, 2.1, 3.8, 1.2, size=27, color=PRIMARY_INK, bold=True)
    add_card(slide, 7.55, 4.08, 4.35, 1.2, fill=RGBColor(39, 39, 42), border=RGBColor(63, 63, 70))
    add_text(slide, "승인된 근거", 7.82, 4.27, 2.2, 0.3, size=10, color=RGBColor(161, 161, 170))
    add_text(slide, "공개 화면에는 검증된 내용만 표시", 7.82, 4.65, 3.7, 0.35, size=11, color=PRIMARY_INK)

    slide = base_slide(deck, "START", 7)
    add_text(slide, "준비됐습니다.", 0.8, 1.35, 8.0, 1.05, size=40, bold=True)
    add_text(slide, "이제 발표에 집중하세요.", 0.82, 2.4, 8.0, 0.7, size=20, color=MUTED)
    add_card(slide, 0.82, 4.25, 4.25, 1.05, fill=ACCENT, border=ACCENT)
    add_text(slide, "발표 시작", 1.15, 4.52, 3.55, 0.48, size=16, color=PRIMARY_INK, bold=True)
    add_text(slide, "impromptu.local", 0.82, 6.55, 3.3, 0.35, size=10, color=MUTED, font=MONO)
    return deck


def add_transitions(path: Path) -> None:
    namespace = "http://schemas.openxmlformats.org/presentationml/2006/main"
    ElementTree.register_namespace("a", "http://schemas.openxmlformats.org/drawingml/2006/main")
    ElementTree.register_namespace("p", namespace)
    with tempfile.TemporaryDirectory() as temp:
        root = Path(temp)
        with zipfile.ZipFile(path) as archive:
            archive.extractall(root)
        slides = sorted((root / "ppt" / "slides").glob("slide*.xml"))
        for index, slide_path in enumerate(slides, start=1):
            tree = ElementTree.parse(slide_path)
            slide = tree.getroot()
            transition = ElementTree.Element(f"{{{namespace}}}transition", {"spd": "med"})
            effect = "push" if index in {3, 5} else "fade"
            attributes = {"dir": "l"} if effect == "push" else {}
            ElementTree.SubElement(transition, f"{{{namespace}}}{effect}", attributes)
            insertion = min(len(slide), 1)
            slide.insert(insertion, transition)
            tree.write(slide_path, encoding="UTF-8", xml_declaration=True)
        rebuilt = root / path.name
        with zipfile.ZipFile(rebuilt, "w", zipfile.ZIP_DEFLATED) as archive:
            for item in sorted(root.rglob("*")):
                if item.is_file() and item != rebuilt:
                    archive.write(item, item.relative_to(root))
        shutil.copyfile(rebuilt, path)


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    deck = build_deck()
    deck.save(str(PPTX_PATH))
    add_transitions(PPTX_PATH)
    subprocess.run(
        [
            "/opt/homebrew/bin/soffice",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(OUTPUT_DIR),
            str(PPTX_PATH),
        ],
        check=True,
        timeout=120,
    )
    if not PDF_PATH.is_file():
        raise FileNotFoundError(PDF_PATH)


if __name__ == "__main__":
    main()
