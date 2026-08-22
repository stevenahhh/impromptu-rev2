# /// script
# requires-python = ">=3.14,<3.15"
# dependencies = ["PyMuPDF==1.28.2", "python-pptx==1.0.2"]
# ///
"""Generate byte-stable Korean PPTX, text-layer PDF, and scanned PDF fixtures.

The fixtures carry a dense six-slide Korean business deck so that indexed
retrieval chunks contain realistic facts (figures, units, years, product and
company names) that the deterministic evidence gate can reconcile verbatim.
The OCR sentinel line ``형식 중립 근거 자료 2026`` is preserved exactly as the
first slide title of every fixture; only body content was added around it.
"""

from __future__ import annotations

import hashlib
import json
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

import pymupdf
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt

ROOT = Path(__file__).resolve().parents[1]
FIXTURE_DIR = ROOT / "tests" / "fixtures" / "format-neutral-decks"
PPTX_PATH = FIXTURE_DIR / "korean-structural.pptx"
TEXT_PDF_PATH = FIXTURE_DIR / "korean-text-layer.pdf"
SCANNED_PDF_PATH = FIXTURE_DIR / "korean-scanned.pdf"
MANIFEST_PATH = FIXTURE_DIR / "fixture-manifest.json"

SENTINEL = "형식 중립 근거 자료 2026"
FORBIDDEN_PPTX_SENTINEL = "<date/time>"
FIXED_DATETIME = datetime(2024, 1, 1, tzinfo=UTC)
FIXED_ZIP_TIME = (2024, 1, 1, 0, 0, 0)
PAGE_WIDTH = 720
PAGE_HEIGHT = 405

TITLE_FONT_SIZE = 24
BODY_FONT_SIZE = 14
TITLE_BASELINE_Y = 70
FIRST_BODY_BASELINE_Y = 120
BODY_LINE_STEP = 36
TEXT_ORIGIN_X = 60


@dataclass(frozen=True)
class Slide:
    """One slide as an ordered list of single-line text strings."""

    title: str
    lines: tuple[str, ...]

    def texts(self) -> tuple[str, ...]:
        return (self.title, *self.lines)


# Dense synthetic Korean business deck. Every line must stay short enough to fit
# PAGE_WIDTH at BODY_FONT_SIZE, use glyphs available in the built-in Korea CJK
# font (Hangul, ASCII digits/letters, basic punctuation), and keep concrete fact
# tokens (numbers, units, years, proper nouns) verbatim for evidence gating.
SLIDES: tuple[Slide, ...] = (
    Slide(
        SENTINEL,
        (
            "한빛유통 2026년 사업 전략 보고",
            "작성: 경영기획본부, 2026년 1월 15일",
        ),
    ),
    Slide(
        "매출 실적 요약",
        (
            "2025년 연간 매출은 482억 원으로 전년 대비 23% 증가했다.",
            "대면 판매 채널 매출은 217억 원이고 온라인 몰 매출은 265억 원이다.",
            "4분기 매출은 분기 사상 최대인 151억 원을 기록했다.",
        ),
    ),
    Slide(
        "구독 사업 현황",
        (
            "유료 구독자 수는 2025년 말 기준 12,400명이다.",
            "월 구독 객단가는 19,000원이고 연 구독 유지율은 78%다.",
            "2026년 말까지 구독자 20,000명 달성을 목표로 한다.",
        ),
    ),
    Slide(
        "신규 상품 계획",
        (
            "2026년 3월에 프리미엄 배송 상품 단비를 출시한다.",
            "단비는 서울과 부산 등 6개 권역에서 우선 운영한다.",
            "초기 투자는 85억 원이고 손익분기 도달 시점은 2028년으로 본다.",
        ),
    ),
    Slide(
        "시장 환경과 경쟁",
        (
            "국내 유통 시장의 온라인 비중은 2025년에 42.7%였다.",
            "경쟁사 모아마켓의 점유율은 18.2%로 업계 2위다.",
            "물류 인건비는 최근 3년간 연평균 9.6% 올랐다.",
        ),
    ),
    Slide(
        "2026년 실행 계획",
        (
            "1분기에 물류센터 자동화 로봇 12대를 도입한다.",
            "2분기에 충청권 신규 물류센터를 준공한다.",
            "3분기에 단비 권역을 12개로 확대하고 앱 3.0을 출시한다.",
            "연말까지 매출 610억 원과 영업이익률 8.5%를 달성한다.",
        ),
    ),
)


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _zip_info(name: str) -> ZipInfo:
    info = ZipInfo(name, FIXED_ZIP_TIME)
    info.compress_type = ZIP_DEFLATED
    info.create_system = 3
    info.external_attr = 0o600 << 16
    return info


def _normalize_pptx_archive(path: Path) -> None:
    with ZipFile(path) as source:
        entries = {name: source.read(name) for name in source.namelist() if not name.endswith("/")}
    with TemporaryDirectory() as temporary:
        rebuilt = Path(temporary) / path.name
        with ZipFile(rebuilt, "w", compression=ZIP_DEFLATED, compresslevel=9) as output:
            for name in sorted(entries):
                output.writestr(_zip_info(name), entries[name])
        path.write_bytes(rebuilt.read_bytes())


def _add_line_shape(slide: object, text: str, baseline_y: float, font_size: float) -> None:
    box = slide.shapes.add_textbox(Pt(TEXT_ORIGIN_X), Pt(baseline_y - font_size), Pt(600), Pt(font_size + 8))  # type: ignore[attr-defined]
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = False
    run = frame.paragraphs[0].add_run()
    run.text = text
    run.font.name = "Apple SD Gothic Neo"
    run.font.size = Pt(font_size)
    run.font.color.rgb = RGBColor(24, 24, 27)


def _generate_pptx() -> None:
    deck = Presentation()
    deck.slide_width = Pt(PAGE_WIDTH)
    deck.slide_height = Pt(PAGE_HEIGHT)
    deck.core_properties.title = "Format-neutral Korean ingestion fixture"
    deck.core_properties.subject = "Synthetic fixture for structural extraction"
    deck.core_properties.author = "Impromptu fixture generator"
    deck.core_properties.created = FIXED_DATETIME
    deck.core_properties.modified = FIXED_DATETIME
    deck.core_properties.last_modified_by = "Impromptu fixture generator"
    deck.core_properties.revision = 1

    for slide_spec in SLIDES:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(250, 250, 250)
        _add_line_shape(slide, slide_spec.title, TITLE_BASELINE_Y, TITLE_FONT_SIZE)
        for offset, line in enumerate(slide_spec.lines):
            _add_line_shape(
                slide,
                line,
                FIRST_BODY_BASELINE_Y + offset * BODY_LINE_STEP,
                BODY_FONT_SIZE,
            )

    deck.save(PPTX_PATH)
    _normalize_pptx_archive(PPTX_PATH)


def _set_pdf_metadata(document: pymupdf.Document, title: str) -> None:
    document.set_metadata(
        {
            "title": title,
            "author": "Impromptu fixture generator",
            "subject": "Synthetic Korean ingestion fixture",
            "keywords": "impromptu,fixture,korean",
            "creator": "scripts/generate-format-neutral-deck-fixtures.py",
            "producer": "PyMuPDF 1.28.2",
            "creationDate": "D:20240101000000Z",
            "modDate": "D:20240101000000Z",
        }
    )


def _save_pdf(document: pymupdf.Document, path: Path) -> None:
    document.save(path, garbage=4, clean=True, deflate=True, no_new_id=True)
    document.close()


def _insert_fixture_text(page: pymupdf.Page, slide_index: int) -> None:
    page.draw_rect(
        pymupdf.Rect(36, 36, PAGE_WIDTH - 36, PAGE_HEIGHT - 36),
        color=(0.145, 0.388, 0.922),
        width=2,
    )
    slide_spec = SLIDES[slide_index]
    page.insert_text(
        (TEXT_ORIGIN_X, TITLE_BASELINE_Y),
        slide_spec.title,
        fontname="korea",
        fontsize=TITLE_FONT_SIZE,
        color=(0.094, 0.094, 0.106),
    )
    for offset, line in enumerate(slide_spec.lines):
        page.insert_text(
            (TEXT_ORIGIN_X, FIRST_BODY_BASELINE_Y + offset * BODY_LINE_STEP),
            line,
            fontname="korea",
            fontsize=BODY_FONT_SIZE,
            color=(0.094, 0.094, 0.106),
        )


def _generate_text_pdf() -> None:
    document = pymupdf.open()
    for index in range(len(SLIDES)):
        page = document.new_page(width=PAGE_WIDTH, height=PAGE_HEIGHT)
        _insert_fixture_text(page, index)
    _set_pdf_metadata(document, "Korean text-layer fixture")
    _save_pdf(document, TEXT_PDF_PATH)


def _generate_scanned_pdf() -> None:
    source = pymupdf.open()
    for index in range(len(SLIDES)):
        source_page = source.new_page(width=PAGE_WIDTH, height=PAGE_HEIGHT)
        _insert_fixture_text(source_page, index)

    document = pymupdf.open()
    for source_page in source:
        pixmap = source_page.get_pixmap(
            matrix=pymupdf.Matrix(2, 2), colorspace=pymupdf.csRGB, alpha=False
        )
        image = pixmap.tobytes("png")
        page = document.new_page(width=PAGE_WIDTH, height=PAGE_HEIGHT)
        page.insert_image(page.rect, stream=image)
    source.close()
    _set_pdf_metadata(document, "Korean scanned-image fixture")
    _save_pdf(document, SCANNED_PDF_PATH)


def _density_block() -> dict[str, object]:
    texts = tuple(text for slide in SLIDES for text in slide.texts())
    return {
        "slideCount": len(SLIDES),
        "textElementCount": len(texts),
        "totalTextCharacters": sum(len(text) for text in texts),
        "note": (
            "Pinned machine-checkable density invariants for the structural "
            "fixtures; the freeze tests assert them against real extraction."
        ),
    }


def _fixture_entry(
    fixture_id: str,
    kind: str,
    path: Path,
    source_method: str,
    expected_semantics: str,
) -> dict[str, object]:
    return {
        "id": fixture_id,
        "kind": kind,
        "path": path.relative_to(ROOT).as_posix(),
        "sha256": _sha256(path),
        "expectedSemantics": expected_semantics,
        "provenance": {
            "origin": "repository-generated synthetic fixture",
            "sourceMethod": source_method,
            "thirdPartySourceMaterial": False,
            "generatedBy": "scripts/generate-format-neutral-deck-fixtures.py",
        },
    }


def _write_manifest() -> None:
    structural_semantics = (
        f"DeckManifest has {len(SLIDES)} slides whose ordered TextElement tuples are identical "
        "across the PPTX and text-layer PDF fixtures; slide one opens with the sentinel"
    )
    manifest = {
        "schemaVersion": 1,
        "registryId": "impromptu-format-neutral-korean-decks-v2",
        "generator": {
            "path": "scripts/generate-format-neutral-deck-fixtures.py",
            "command": "uv run scripts/generate-format-neutral-deck-fixtures.py",
            "requiresPython": ">=3.14,<3.15",
            "dependencies": ["PyMuPDF==1.28.2", "python-pptx==1.0.2"],
        },
        "sentinels": {
            "structural": SENTINEL,
            "ocr": SENTINEL,
            "forbiddenPptxText": FORBIDDEN_PPTX_SENTINEL,
        },
        "density": _density_block(),
        "fixtures": [
            _fixture_entry(
                "korean-structural-pptx",
                "pptx",
                PPTX_PATH,
                "python-pptx six-slide deck with one native Korean textbox per line; "
                "ZIP metadata normalized",
                structural_semantics,
            ),
            _fixture_entry(
                "korean-text-layer-pdf",
                "text-layer-pdf",
                TEXT_PDF_PATH,
                "PyMuPDF pages with native Korean PDF text using the built-in Korea CJK font",
                structural_semantics,
            ),
            _fixture_entry(
                "korean-scanned-pdf",
                "scanned-pdf",
                SCANNED_PDF_PATH,
                "2x rasterization of the same six synthetic Korean pages embedded as each "
                "PDF page's only image",
                "DeckManifest has no TextElement and warns scanned_page_requires_ocr per page; "
                "OCR must restore the sentinel line exactly on slide one",
            ),
        ],
        "migrationReservations": {
            "private": [
                {"filename": "0008_deck_retrieval_hybrid.sql", "owner": "retrieval"},
                {"filename": "0009_session_reports.sql", "owner": "report"},
            ],
            "projection": {"latest": "0005_gateway_state.sql", "newMigrationAllowed": False},
        },
    }
    serialized = json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True) + "\n"
    # Match the repository's Biome format for short arrays so a regenerated
    # manifest is byte-identical to the committed one.
    serialized = serialized.replace(
        '''    "dependencies": [
      "PyMuPDF==1.28.2",
      "python-pptx==1.0.2"
    ]''',
        '    "dependencies": ["PyMuPDF==1.28.2", "python-pptx==1.0.2"]',
    )
    MANIFEST_PATH.write_text(serialized, encoding="utf-8")


def main() -> None:
    FIXTURE_DIR.mkdir(parents=True, exist_ok=True)
    _generate_pptx()
    _generate_text_pdf()
    _generate_scanned_pdf()
    _write_manifest()
    for fixture in (PPTX_PATH, TEXT_PDF_PATH, SCANNED_PDF_PATH):
        print(f"{fixture.relative_to(ROOT)} {_sha256(fixture)}")


if __name__ == "__main__":
    main()
