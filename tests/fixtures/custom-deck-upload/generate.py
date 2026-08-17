"""Generate byte-stable custom-deck upload E2E fixtures.

The PPTX carries three real OOXML click groups (entrance, emphasis, exit), a
fade transition, and ECMA-376 obfuscated embedded TrueType data. The PDF has
three pages with distinct dimensions. Zip timestamps and document metadata are
fixed so rerunning this file produces identical bytes.
"""

from __future__ import annotations

import hashlib
import os
from datetime import UTC, datetime
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

import pymupdf
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

ROOT = Path(__file__).parent
PPTX_PATH = ROOT / "custom-runtime-deck.pptx"
PDF_PATH = ROOT / "custom-static-deck.pdf"
FIXED_TIME = (2024, 1, 1, 0, 0, 0)
FONT_GUID = "{001B70DC-AA60-4AD5-90EC-18A0948E1EAE}"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CONTENT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
FONT_FAMILY = "DejaVu Sans"
FONT_SHA256 = "7da195a74c55bef988d0d48f9508bd5d849425c1770dba5d7bfc6ce9ed848954"
DEFAULT_FONT_PATH = Path.home() / "Library/Fonts/DejaVuSans.ttf"


def _sfnt_font_bytes() -> bytes:
    """Read the pinned DejaVu Sans 2.37 TTF installed by Homebrew's font-dejavu cask."""
    font_path = Path(os.environ.get("IMPROMPTU_E2E_FONT_PATH", DEFAULT_FONT_PATH))
    try:
        payload = font_path.read_bytes()
    except OSError as error:
        raise RuntimeError(
            f"E2E font missing at {font_path}; install Homebrew cask font-dejavu 2.37"
        ) from error
    digest = hashlib.sha256(payload).hexdigest()
    if digest != FONT_SHA256:
        raise RuntimeError(f"unexpected DejaVu Sans bytes at {font_path}: {digest}")
    return payload


def _obfuscate_font(font: bytes) -> bytes:
    key = bytes.fromhex(FONT_GUID.strip("{}").replace("-", ""))[::-1]
    result = bytearray(font)
    for index in range(min(32, len(result))):
        result[index] ^= key[index % 16]
    return bytes(result)


def _behavior(
    shape_id: int, effect_class: str, direction: str, color: bool = False
) -> str:
    if color:
        behavior = f"""<p:anim clrSpc="rgb" from="#0057B8" to="#F15A24">
          <p:cBhvr><p:cTn dur="240"/><p:tgtEl><p:spTgt spid="{shape_id}"/></p:tgtEl></p:cBhvr>
        </p:anim>"""
    else:
        behavior = f"""<p:animEffect transition="{direction}" filter="fade">
          <p:cBhvr><p:cTn dur="180"/><p:tgtEl><p:spTgt spid="{shape_id}"/></p:tgtEl></p:cBhvr>
        </p:animEffect>"""
    preset = {"entr": 1, "emph": 11, "exit": 10}[effect_class]
    return f"""<p:par><p:cTn presetID="{preset}" presetClass="{effect_class}"
      presetSubtype="0" nodeType="clickEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst>
      <p:childTnLst>{behavior}</p:childTnLst></p:cTn></p:par>"""


def _timing(shape_ids: tuple[int, int, int]) -> etree._Element:
    groups = (
        _behavior(shape_ids[0], "entr", "in"),
        _behavior(shape_ids[1], "emph", "in", color=True),
        _behavior(shape_ids[2], "exit", "out"),
    )
    wrapped_groups = "".join(
        f"<p:par><p:cTn><p:childTnLst>{group}</p:childTnLst></p:cTn></p:par>"
        for group in groups
    )
    xml = f"""<p:timing xmlns:p="{P_NS}"><p:tnLst><p:par><p:cTn nodeType="tmRoot">
      <p:childTnLst><p:seq><p:cTn nodeType="mainSeq"><p:childTnLst>{wrapped_groups}
      </p:childTnLst></p:cTn></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>"""
    return etree.fromstring(xml.encode())


def _zip_info(name: str) -> ZipInfo:
    info = ZipInfo(name, FIXED_TIME)
    info.compress_type = ZIP_DEFLATED
    info.external_attr = 0o600 << 16
    return info


def _generate_pptx() -> None:
    with TemporaryDirectory() as temporary:
        original = Path(temporary) / "original.pptx"
        deck = Presentation()
        deck.slide_width = Inches(13.333333)
        deck.slide_height = Inches(7.5)
        fixed = datetime(2024, 1, 1, tzinfo=UTC)
        deck.core_properties.created = fixed
        deck.core_properties.modified = fixed
        deck.core_properties.author = "Impromptu QA"
        deck.core_properties.title = "Custom runtime deck"
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(0xF4, 0xF1, 0xEA)
        specs = (
            ("Entrance target", 1.0, 1.0, 3.0, 1.25, RGBColor(0x00, 0x57, 0xB8)),
            ("Emphasis target", 5.15, 2.6, 3.0, 1.25, RGBColor(0x7A, 0x3D, 0xB8)),
            ("Exit target", 9.3, 4.5, 3.0, 1.25, RGBColor(0xF1, 0x5A, 0x24)),
        )
        shapes = []
        for name, left, top, width, height, color in specs:
            shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(left),
                Inches(top),
                Inches(width),
                Inches(height),
            )
            shape.name = name
            shape.fill.solid()
            shape.fill.fore_color.rgb = color
            shape.line.color.rgb = RGBColor(0x24, 0x2A, 0x34)
            shape.line.width = Pt(2)
            shape.text_frame.text = name
            run = shape.text_frame.paragraphs[0].runs[0]
            run.font.name = FONT_FAMILY
            run.font.size = Pt(24)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            shapes.append(shape)
        deck.save(original)

        with ZipFile(original) as source:
            entries = {name: source.read(name) for name in source.namelist()}

        slide_root = etree.fromstring(entries["ppt/slides/slide1.xml"])
        slide_root.append(
            etree.fromstring(
                f'<p:transition xmlns:p="{P_NS}" advClick="1"><p:fade/></p:transition>'.encode()
            )
        )
        slide_root.append(_timing(tuple(shape.shape_id for shape in shapes)))
        entries["ppt/slides/slide1.xml"] = etree.tostring(
            slide_root, xml_declaration=True, encoding="UTF-8", standalone=True
        )

        presentation = etree.fromstring(entries["ppt/presentation.xml"])
        embedded = etree.fromstring(
            f"""<p:embeddedFontLst xmlns:p="{P_NS}" xmlns:r="{R_NS}">
              <p:embeddedFont><p:font typeface="{FONT_FAMILY}"/>
              <p:regular r:id="rIdImpromptuFont"/></p:embeddedFont></p:embeddedFontLst>""".encode()
        )
        presentation.insert(max(0, len(presentation) - 1), embedded)
        entries["ppt/presentation.xml"] = etree.tostring(
            presentation, xml_declaration=True, encoding="UTF-8", standalone=True
        )

        relationships = etree.fromstring(entries["ppt/_rels/presentation.xml.rels"])
        relation = etree.Element(f"{{{PKG_REL_NS}}}Relationship")
        relation.attrib.update(
            {
                "Id": "rIdImpromptuFont",
                "Type": f"{R_NS}/font",
                "Target": "fonts/impromptu.fntdata",
                "Guid": FONT_GUID,
            }
        )
        relationships.append(relation)
        entries["ppt/_rels/presentation.xml.rels"] = etree.tostring(
            relationships, xml_declaration=True, encoding="UTF-8", standalone=True
        )

        content_types = etree.fromstring(entries["[Content_Types].xml"])
        if not any(node.get("Extension") == "fntdata" for node in content_types):
            default = etree.Element(f"{{{CONTENT_NS}}}Default")
            default.attrib.update(
                {"Extension": "fntdata", "ContentType": "application/x-fontdata"}
            )
            content_types.insert(0, default)
        entries["[Content_Types].xml"] = etree.tostring(
            content_types, xml_declaration=True, encoding="UTF-8", standalone=True
        )
        entries["ppt/fonts/impromptu.fntdata"] = _obfuscate_font(_sfnt_font_bytes())

        with ZipFile(PPTX_PATH, "w") as output:
            for name in sorted(entries):
                output.writestr(_zip_info(name), entries[name])


def _generate_pdf() -> None:
    document = pymupdf.open()
    pages = (
        (720, 405, "Landscape page 1"),
        (612, 792, "Portrait page 2"),
        (800, 600, "Static page 3"),
    )
    for width, height, label in pages:
        page = document.new_page(width=width, height=height)
        page.draw_rect(
            pymupdf.Rect(36, 36, width - 36, height - 36),
            color=(0, 0.34, 0.72),
            width=3,
        )
        page.insert_text((72, 96), label, fontsize=24, color=(0.05, 0.08, 0.12))
        page.insert_text((72, 136), "Deterministic custom deck upload QA", fontsize=14)
    document.set_metadata(
        {
            "title": "Custom static deck",
            "author": "Impromptu QA",
            "creationDate": "D:20240101000000Z",
            "modDate": "D:20240101000000Z",
        }
    )
    document.save(PDF_PATH, garbage=4, deflate=True, no_new_id=True)
    document.close()


if __name__ == "__main__":
    ROOT.mkdir(parents=True, exist_ok=True)
    _generate_pptx()
    _generate_pdf()
    for fixture in (PPTX_PATH, PDF_PATH):
        print(f"{fixture.name} {hashlib.sha256(fixture.read_bytes()).hexdigest()}")
